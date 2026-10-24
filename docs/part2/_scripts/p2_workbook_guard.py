"""Publish only audited Artifact Tool cell edits, preserving the original package."""
from pathlib import Path
from copy import deepcopy
import hashlib
import json
import re
import shutil
import sys
import zipfile
import xml.etree.ElementTree as ET
import openpyxl

ROOT=Path(r'D:\RAD\FINAL\Part2'); SCRATCH=ROOT/'notes/P2-root'
NAME='Team02_RAD_Sem01_2627_Team_Workbook.xlsx'
NS='http://schemas.openxmlformats.org/spreadsheetml/2006/main'; N={'m':NS}
ET.register_namespace('',NS)
def tag(t): return '{'+NS+'}'+t
def cells(t): return {c.get('r'):c for c in t.findall('m:sheetData/m:row/m:c',N)}
def colnum(a):
    n=0
    for c in re.match('[A-Z]+',a)[0]: n=n*26+ord(c)-64
    return n
def sha(p): return hashlib.sha256(p.read_bytes()).hexdigest()
def conditional_rules(sheet):
    return [(str(cf.sqref),[ET.tostring(rule.to_tree(),encoding='unicode') for rule in sheet.conditional_formatting[cf]]) for cf in sheet.conditional_formatting]

def main():
    source=SCRATCH/'workbook_before.xlsx'; tool=SCRATCH/'workbook_artifact_roundtrip.xlsx'
    target=SCRATCH/'workbook_candidate.xlsx'; working=ROOT/'workbook'/NAME
    assert sha(source)==sha(working),'Working copy changed after snapshot; re-audit before publishing'
    updates=json.loads((ROOT/'notes/P2_workbook_updates.json').read_text(encoding='utf-8'))
    layout=json.loads((ROOT/'notes/P2_workbook_layout.json').read_text(encoding='utf-8'))
    old=openpyxl.load_workbook(source,data_only=False); allowed={}
    old_cache=openpyxl.load_workbook(source,data_only=True)
    for u in updates:
        assert u['sheet'] in ('Events Table','CRUD Check','Team & Use Cases')
        c=old[u['sheet']][u['cell']]
        assert c.data_type!='f' and c.fill.fgColor.rgb!='FFDDF0DF',(u,'protected')
        allowed.setdefault(u['sheet'],{})[u['cell']]=u['value']
    with zipfile.ZipFile(source) as zin,zipfile.ZipFile(tool) as ztool:
        payload={n:zin.read(n) for n in zin.namelist()}; ss=[]
        if 'xl/sharedStrings.xml' in ztool.namelist():
            ss=[''.join(x.itertext()) for x in ET.fromstring(ztool.read('xl/sharedStrings.xml'))]
        changed_parts=[]
        for idx,sheet in enumerate(old.sheetnames,1):
            if sheet not in allowed: continue
            path=f'xl/worksheets/sheet{idx}.xml'; tree=ET.fromstring(payload[path]); dst=cells(tree)
            authored=cells(ET.fromstring(ztool.read(path))); sd=tree.find('m:sheetData',N)
            rows={int(r.get('r')):r for r in sd}
            for address in allowed[sheet]:
                c=deepcopy(authored.get(address,ET.Element(tag('c'),{'r':address})))
                if c.get('t')=='s':
                    value=ss[int(c.find('m:v',N).text)]
                    for child in list(c): c.remove(child)
                    c.set('t','inlineStr'); ET.SubElement(ET.SubElement(c,tag('is')),tag('t')).text=value
                c.attrib.pop('s',None)
                if address in dst and dst[address].get('s') is not None: c.set('s',dst[address].get('s'))
                num=int(re.search(r'\d+',address)[0])
                assert num in rows,'No new rows in Phase 2'
                if address in dst: rows[num].remove(dst[address])
                rows[num].append(c); dst[address]=c
            for a,c in dst.items():
                if c.find('m:f',N) is not None:
                    assert a in authored,(sheet,a,'missing authored formula')
                    if authored[a].find('m:v',N) is None:
                        # Export omits a cache for IF(...,"",...) empty results.
                        # Only permit previously empty, still-unused template rows.
                        num=int(re.search(r'\d+',a)[0])
                        assert old_cache[sheet][a].value in (None,''),(sheet,a,'nonempty cache lost')
                        assert (sheet=='CRUD Check' and num>=18) or (sheet=='Team & Use Cases' and num in (21,22)),(sheet,a,'unexpected missing cache')
                        continue
                    for v in c.findall('m:v',N): c.remove(v)
                    c.append(deepcopy(authored[a].find('m:v',N)))
                    if authored[a].get('t'): c.set('t',authored[a].get('t'))
                    else: c.attrib.pop('t',None)
            for num,height in layout.get(sheet,{}).items():
                rows[int(num)].set('ht',str(height)); rows[int(num)].set('customHeight','1')
            for row in rows.values(): row[:]=sorted(row,key=lambda c:colnum(c.get('r')))
            payload[path]=ET.tostring(tree,encoding='utf-8',xml_declaration=True); changed_parts.append(path)
        with zipfile.ZipFile(target,'w',zipfile.ZIP_DEFLATED) as zout:
            for n,b in payload.items(): zout.writestr(n,b)
        assert all(payload[n]==zin.read(n) for n in payload if n not in changed_parts),'Unrelated XLSX part changed'
    new=openpyxl.load_workbook(target,data_only=False)
    check={'formulas':0,'lecturer_cells':0,'validation_rules':0,'edited_cells':len(updates),'changed_parts':changed_parts}
    assert old.sheetnames==new.sheetnames
    for sheet in old.sheetnames:
        a,b=old[sheet],new[sheet]
        assert str(a.data_validations)==str(b.data_validations),(sheet,'validations')
        check['validation_rules']+=len(a.data_validations.dataValidation)
        assert conditional_rules(a)==conditional_rules(b),(sheet,'conditional formatting')
        assert list(a.merged_cells.ranges)==list(b.merged_cells.ranges)
        for row in a:
            for c in row:
                n=b[c.coordinate]; expected=allowed.get(sheet,{}).get(c.coordinate,c.value)
                assert n.value==expected,(sheet,c.coordinate,'content',n.value,expected)
                assert c._style==n._style,(sheet,c.coordinate,'style')
                if c.data_type=='f': assert n.data_type=='f'; check['formulas']+=1
                if c.fill.fgColor.rgb=='FFDDF0DF': assert c.value==n.value; check['lecturer_cells']+=1
    cache=openpyxl.load_workbook(target,data_only=True)
    check['crud_gap_values']=[cache['CRUD Check'][f'G{r}'].value for r in range(5,18)]
    assert check['crud_gap_values']==['OK']*13
    check['client_count']=cache['Client Q&A']['C16'].value; assert check['client_count']==5
    check['preservation_pass']=True; check['baseline_sha256']=sha(source); check['candidate_sha256']=sha(target)
    # Publication is after every audit passes and after the Wave 2 barrier.
    audit_only='--audit-only' in sys.argv
    check['published']=not audit_only
    if not audit_only:
        shutil.copyfile(target,working); assert sha(working)==sha(target)
    output='P2_workbook_candidate_audit.json' if audit_only else 'P2_workbook_preservation.json'
    (ROOT/'notes'/output).write_text(json.dumps(check,indent=2),encoding='utf-8')
    print(json.dumps(check,indent=2))

if __name__=='__main__': main()
