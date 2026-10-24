"""Read primary authorities and preserve the Phase 2 immutable boundaries."""
from pathlib import Path
import ast
import hashlib
import json
import subprocess
import sys
import zipfile
from docx import Document
import openpyxl
from pypdf import PdfReader

ROOT = Path(r'D:\RAD\FINAL')
REPO = ROOT / 'SaigonBite'
NOTES = ROOT / 'Part2/notes'
NOTES.mkdir(parents=True, exist_ok=True)

def sha(p):
    return hashlib.sha256(p.read_bytes()).hexdigest()

def manifest():
    data = {}
    for label, base in [('part1',ROOT/'Part1_v2'),('repo_part1',REPO/'docs/part1'),
                        ('app',REPO/'saigonbite'),('tests',REPO/'tests'),('db',REPO/'db')]:
        data[label] = {p.relative_to(base).as_posix():sha(p) for p in sorted(base.rglob('*'))
                       if p.is_file() and not any(x in p.parts for x in ('__pycache__','.pytest_cache'))}
    data['entrypoints'] = {n:sha(REPO/n) for n in ('app.py','init_db.py')}
    data['contracts'] = {n:sha(ROOT/n) for n in ('PLAN_SaigonBite_Part2_Final.md','Part2/content/links.json')}
    data['section_I'] = {p.name:sha(p) for p in sorted((REPO/'docs/part2/report').glob('I_*.md'))}
    return data

def main():
    if len(sys.argv)>1 and sys.argv[1]=='end':
        start=json.loads((NOTES/'P2_start_hashes.json').read_text(encoding='utf-8'))
        now=manifest()
        changes={k:[n for n in sorted(set(start[k])|set(v)) if start[k].get(n)!=v.get(n)] for k,v in now.items()}
        diff=subprocess.run(['git','diff','--exit-code','657089b','--','saigonbite','tests','db','app.py','init_db.py'],cwd=REPO,capture_output=True,text=True)
        out={'counts':{k:len(v) for k,v in now.items()},'changes':changes,'git_app_diff_exit':diff.returncode,'pass':not any(changes.values()) and diff.returncode==0}
        (NOTES/'P2_end_hash_check.json').write_text(json.dumps(out,indent=2),encoding='utf-8')
        print(json.dumps(out,indent=2)); assert out['pass']; return
    status=subprocess.check_output(['git','status','--porcelain'],cwd=REPO,text=True)
    branch=subprocess.check_output(['git','branch','--show-current'],cwd=REPO,text=True).strip()
    head=subprocess.check_output(['git','rev-parse','HEAD'],cwd=REPO,text=True).strip()
    # This script is the only authorized untracked preflight source.
    assert not status.replace('?? docs/part2/_scripts/p2_preflight.py\n','').strip(),status
    assert branch=='part2'
    subprocess.run(['git','merge-base','--is-ancestor','657089b','HEAD'],cwd=REPO,check=True)
    target=NOTES/'P2_start_hashes.json'; assert not target.exists(),'Never overwrite start manifest'
    target.write_text(json.dumps(manifest(),indent=2),encoding='utf-8')
    out={'branch':branch,'head':head,'preflight_only_untracked':status,'primary':{}}
    doc=Document(ROOT/'Lec_re/RAD_Sem01_2627_Report_Template.docx')
    out['primary']['template_paragraphs']=[p.text for p in doc.paragraphs if p.text.strip()]
    out['primary']['template_tables']=[[[c.text for c in r.cells] for r in t.rows] for t in doc.tables]
    for name in ['RAD_Sem01_2627_MarkingSheet.xlsx','RAD_Sem01_2627_Team_Workbook_Template.xlsx']:
        wb=openpyxl.load_workbook(ROOT/'Lec_re'/name,data_only=False)
        out['primary'][name]={s.title:[[c.value for c in r] for r in s] for s in wb}
    wb=openpyxl.load_workbook(ROOT/'Part2/workbook/Team02_RAD_Sem01_2627_Team_Workbook.xlsx',data_only=False)
    out['working_sheets']={n:[[c.value for c in r] for r in wb[n]] for n in ['Team & Use Cases','Events Table','CRUD Check']}
    for name in ['RAD_Sem01_2627_FinalProject_Requirements_SaigonBite.pdf','RAD_Sem01_2627_Case_Study_SaigonBite.pdf']:
        out['primary'][name]=[p.extract_text() for p in PdfReader(ROOT/'Lec_re'/name).pages]
    out['controllers']={}
    for p in sorted((REPO/'saigonbite/business/controllers').glob('*.py')):
        tree=ast.parse(p.read_text(encoding='utf-8-sig'))
        out['controllers'][p.name]={n.name:[{'method':m.name,'line':m.lineno,'parameters':ast.unparse(m.args)} for m in n.body if isinstance(m,(ast.FunctionDef,ast.AsyncFunctionDef))] for n in tree.body if isinstance(n,ast.ClassDef)}
    with zipfile.ZipFile(ROOT/'balanced-diagrams.skill.txt') as z:
        out['diagram_skill']=z.read('balanced-diagrams/SKILL.md').decode('utf-8')
    (NOTES/'P2_primary_sources.json').write_text(json.dumps(out,indent=2,ensure_ascii=False,default=str),encoding='utf-8')
    print(json.dumps({'branch':branch,'head':head,'counts':{k:len(v) for k,v in manifest().items()},'primary_evidence':'Part2/notes/P2_primary_sources.json'},indent=2))
    print('TEMPLATE TABLES:')
    for t in out['primary']['template_tables']: print(json.dumps(t,ensure_ascii=False))
    for s,rows in out['primary']['RAD_Sem01_2627_MarkingSheet.xlsx'].items():
        for i,row in enumerate(rows,1):
            if (s=='Checklist' and len(row)>1 and row[1] in range(14,21)) or any('P2-II' in str(c) for c in row): print(s,i,json.dumps(row,ensure_ascii=False))

if __name__=='__main__': main()
