"""Preserve Phase 3 boundaries and directly extract local primary authorities.
Author: 722I0004 · AI-assisted (Codex). Never imports application modules.
"""
from pathlib import Path
from datetime import datetime, timezone
import hashlib, json, subprocess, sys, zipfile

ROOT = Path('D:/RAD/FINAL')
REPO = ROOT/'SaigonBite'
NOTES = ROOT/'Part2/notes'

def sha(p): return hashlib.sha256(p.read_bytes()).hexdigest()
def git(*args): return subprocess.check_output(['git',*args],cwd=REPO,text=True,encoding='utf-8').strip()
def files(base):
    return {p.relative_to(base).as_posix():sha(p) for p in sorted(base.rglob('*')) if p.is_file() and not any(x in p.parts for x in ('__pycache__','.pytest_cache'))}
def manifest():
    result={label:files(base) for label,base in [('part1',ROOT/'Part1_v2'),('repo_part1',REPO/'docs/part1'),('app',REPO/'saigonbite'),('tests',REPO/'tests'),('db',REPO/'db')]}
    result['entrypoints']={n:sha(REPO/n) for n in ('app.py','init_db.py')}
    result['contracts']={n:sha(ROOT/n) for n in ('PLAN_SaigonBite_Part2_Final.md','Part2/content/links.json')}
    result['workbook']={str(p.relative_to(ROOT)):sha(p) for base in [ROOT/'Part2/workbook',ROOT/'Part2/drive_upload'] for p in sorted(base.glob('*.xlsx'))}
    result['section_I_II']={p.relative_to(REPO).as_posix():sha(p) for pattern in ('I_*.md','II_*.md') for p in sorted((REPO/'docs/part2/report').glob(pattern))}
    result['phase2_models']=files(REPO/'models')
    result['phase2_contracts']={p.relative_to(REPO).as_posix():sha(p) for p in (REPO/'docs/part2').glob('UC-*/description.md')}
    for name in ('ssd_operations.csv','ssd_flow_membership.csv','pc_mg_register.csv','analysis_consistency.csv'):
        p=REPO/'docs/part2'/name; result['phase2_contracts'][p.relative_to(REPO).as_posix()]=sha(p)
    return result

def main():
    NOTES.mkdir(parents=True,exist_ok=True)
    target=NOTES/'P3_start_hashes.json'
    if len(sys.argv)>1 and sys.argv[1]=='end':
        start=json.loads(target.read_text(encoding='utf-8')); current=manifest(); differences={}
        for group,old in start['protected'].items():
            now=current[group]
            # The existing model bytes are frozen; P3 adds new stems only.
            names=set(old) if group=='phase2_models' else set(old)|set(now)
            differences[group]=[n for n in sorted(names) if old.get(n)!=now.get(n)]
        app_diff=git('diff',start['C0'],'--','saigonbite','tests','db','app.py','init_db.py')
        # In addition to the byte boundaries, protect every pre-existing tracked
        # source (including old agent reports/scripts). New P3 paths are allowed.
        # Git handles the repository's working-tree line-ending normalization;
        # this supplements, and never replaces, the binary start/end checks.
        prior_deltas=git('diff','--diff-filter=MDRT','--name-only',start['C0']).splitlines()
        allowed={'docs/part2/ai_usage_log.md','docs/part2/traceability.csv'}
        unexpected=[name for name in prior_deltas if name not in allowed]
        out={'C0':start['C0'],'HEAD':git('rev-parse','HEAD'),'counts':{k:len(v) for k,v in start['protected'].items()},'changes':differences,'app_diff_empty':not app_diff,'preexisting_repo_deltas':prior_deltas,'unexpected_preexisting_repo_deltas':unexpected,'pass':not any(differences.values()) and not app_diff and not unexpected}
        (NOTES/'P3_end_hash_check.json').write_text(json.dumps(out,indent=2),encoding='utf-8'); print(json.dumps(out,indent=2));assert out['pass'];return
    assert not target.exists(),'Start manifest already exists; never overwrite it'
    assert git('branch','--show-current')=='part2'
    c0=git('rev-parse','HEAD')
    start={'author':'722I0004 · AI-assisted (Codex)','started_at_utc':datetime.now(timezone.utc).isoformat(),'C0':c0,'branch':'part2','status':git('status','--porcelain'),'committed_app_changes_since_G2':git('diff','45ac762',c0,'--','saigonbite','tests','db','app.py','init_db.py'),'uncommitted_app_changes':git('status','--porcelain','--','saigonbite','tests','db','app.py','init_db.py'),'protected':manifest()}
    target.write_text(json.dumps(start,indent=2),encoding='utf-8')
    from docx import Document
    from pypdf import PdfReader
    import openpyxl
    doc=Document(ROOT/'Lec_re/RAD_Sem01_2627_Report_Template.docx')
    primary={'template_paragraphs':[p.text for p in doc.paragraphs if p.text.strip()],'template_tables':[[[c.text for c in row.cells] for row in table.rows] for table in doc.tables]}
    for name in ('RAD_Sem01_2627_MarkingSheet.xlsx','RAD_Sem01_2627_Team_Workbook_Template.xlsx'):
        wb=openpyxl.load_workbook(ROOT/'Lec_re'/name,data_only=False)
        primary[name]={s.title:[[c.value for c in row] for row in s] for s in wb}
    for name in ('RAD_Sem01_2627_FinalProject_Requirements_SaigonBite.pdf','RAD_Sem01_2627_Case_Study_SaigonBite.pdf'):
        primary[name]=[p.extract_text() for p in PdfReader(ROOT/'Lec_re'/name).pages]
    with zipfile.ZipFile(ROOT/'balanced-diagrams.skill.txt') as archive:
        primary['diagram_skill']=archive.read('balanced-diagrams/SKILL.md').decode('utf-8')
    (NOTES/'P3_primary_sources.json').write_text(json.dumps(primary,indent=2,ensure_ascii=False,default=str),encoding='utf-8')
    print(json.dumps({k:v for k,v in start.items() if k!='protected'},indent=2)); print('Protected files:',sum(map(len,start['protected'].values())))

if __name__=='__main__':main()
