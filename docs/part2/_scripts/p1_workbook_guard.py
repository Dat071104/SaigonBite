"""Apply Artifact Tool-authored cells while retaining the original XLSX package.

The Artifact Tool owns cell values and formula calculation. This narrow OOXML
merge retains original styles, validation, lecturer cells and unrelated parts.
openpyxl is used for read-only audits, never workbook authoring.
"""
from pathlib import Path
from copy import deepcopy
import json
import re
import zipfile
import xml.etree.ElementTree as ET
import openpyxl

ROOT = Path(r"D:\RAD\FINAL")
BASE = ROOT / "Part2"
NAME = "Team02_RAD_Sem01_2627_Team_Workbook.xlsx"
NS = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
ET.register_namespace("", NS)
N = {"m": NS}


def tag(name):
    return f"{{{NS}}}{name}"


def cell_index(tree):
    return {c.attrib["r"]: c for c in tree.findall("m:sheetData/m:row/m:c", N)}


def col_num(address):
    result = 0
    for char in re.match(r"[A-Z]+", address)[0]:
        result = result * 26 + ord(char) - 64
    return result


def inline_cell(cell, strings):
    c = deepcopy(cell)
    if c.get("t") == "s":
        text = strings[int(c.find("m:v", N).text)]
        for child in list(c):
            c.remove(child)
        c.set("t", "inlineStr")
        ET.SubElement(ET.SubElement(c, tag("is")), tag("t")).text = text
    return c


def add_validations(tree, specs):
    """Extend only new input rows; retain every existing validation element."""
    if not specs:
        return
    rules = tree.find("m:dataValidations", N)
    if rules is None:
        raise ValueError("Expected the original validation container")
    originals = list(rules)
    for spec in specs:
        source = next(r for r in originals if r.get("sqref") == spec["from"])
        rule = deepcopy(source)
        rule.set("sqref", spec["to"])
        rules.append(rule)
    rules.set("count", str(len(rules)))


def main():
    original = ROOT / "Part1_v2/drive_upload" / NAME
    tool = BASE / "notes/P1_workbook_artifact_roundtrip.xlsx"
    changes = json.loads((BASE / "notes/P1_workbook_updates.json").read_text(encoding="utf-8"))
    layout = json.loads((BASE / "notes/P1_workbook_layout.json").read_text(encoding="utf-8"))
    old = openpyxl.load_workbook(original, data_only=False)
    by_sheet = {}
    for update in changes:
        cell = old[update["sheet"]][update["cell"]]
        if cell.fill.fgColor.rgb == "FFDDF0DF":
            raise ValueError(f"Lecturer cell is protected: {update}")
        if cell.data_type == "f":
            raise ValueError(f"Existing formula cannot be overwritten: {update}")
        by_sheet.setdefault(update["sheet"], []).append(update)
    with zipfile.ZipFile(original) as zin, zipfile.ZipFile(tool) as ztool:
        ss = []
        if "xl/sharedStrings.xml" in ztool.namelist():
            ss = ["".join(si.itertext()) for si in ET.fromstring(ztool.read("xl/sharedStrings.xml"))]
        payload = {name: zin.read(name) for name in zin.namelist()}
        for idx, sheet in enumerate(old.sheetnames, 1):
            path = f"xl/worksheets/sheet{idx}.xml"
            tree = ET.fromstring(payload[path])
            authored = cell_index(ET.fromstring(ztool.read(path)))
            cells = cell_index(tree)
            sd = tree.find("m:sheetData", N)
            rows = {int(r.get("r")): r for r in sd}
            for update in by_sheet.get(sheet, []):
                address = update["cell"]
                # An authored null can be represented by an absent cell.
                c = inline_cell(authored.get(address, ET.Element(tag("c"), {"r": address})), ss)
                source_style = cells.get(address)
                if update.get("format_from"):
                    source_style = cells[update["format_from"]]
                if source_style is None and update.get("style_from"):
                    source_style = cells.get(update["style_from"])
                c.attrib.pop("s", None)
                if source_style is not None and source_style.get("s") is not None:
                    c.set("s", source_style.get("s"))
                rnum = int(re.search(r"\d+", address)[0])
                if rnum not in rows:
                    rows[rnum] = ET.SubElement(sd, tag("row"), {"r": str(rnum)})
                if address in cells:
                    rows[rnum].remove(cells[address])
                rows[rnum].append(c)
                cells[address] = c
            # Keep each original formula exactly; refresh only its tool-calculated cache.
            for address, c in cells.items():
                if c.find("m:f", N) is not None and address in authored:
                    ac = authored[address]
                    for v in c.findall("m:v", N):
                        c.remove(v)
                    if ac.find("m:v", N) is not None:
                        c.append(deepcopy(ac.find("m:v", N)))
                    if ac.get("t"):
                        c.set("t", ac.get("t"))
                    else:
                        c.attrib.pop("t", None)
            spec = layout.get(sheet, {})
            for row, height in spec.get("heights", {}).items():
                if int(row) in rows:
                    rows[int(row)].set("ht", str(height))
                    rows[int(row)].set("customHeight", "1")
            cols = tree.find("m:cols", N)
            for col, width in spec.get("new_columns", {}).items():
                ET.SubElement(cols, tag("col"), {"min": str(col), "max": str(col), "width": str(width), "customWidth": "1"})
            for row in rows.values():
                row[:] = sorted(row, key=lambda c: col_num(c.get("r")))
            sd[:] = [rows[r] for r in sorted(rows)]
            if spec.get("dimension"):
                tree.find("m:dimension", N).set("ref", spec["dimension"])
            for attr in ("autoFilter",):
                element = tree.find(f"m:{attr}", N)
                if element is not None and spec.get(attr):
                    element.set("ref", spec[attr])
            add_validations(tree, spec.get("validations", []))
            if sheet in by_sheet:
                payload[path] = ET.tostring(tree, encoding="utf-8", xml_declaration=True)
    target = BASE / "notes/P1_workbook_candidate.xlsx"
    with zipfile.ZipFile(target, "w", zipfile.ZIP_DEFLATED) as out:
        for name, content in payload.items():
            out.writestr(name, content)
    old = openpyxl.load_workbook(original, data_only=False)
    new = openpyxl.load_workbook(target, data_only=False)
    checks = {"original_formulas": 0, "lecturer_cells": 0, "validation_rules": 0}
    for sheet in old.sheetnames:
        a, b = old[sheet], new[sheet]
        allowed = {u["cell"] for u in by_sheet.get(sheet, [])}
        original_rules = list(a.data_validations.dataValidation)
        actual_rules = list(b.data_validations.dataValidation)
        assert len(actual_rules) >= len(original_rules), sheet
        for expected, actual_rule in zip(original_rules, actual_rules):
            assert str(expected) == str(actual_rule), (sheet, "original validation")
        checks["validation_rules"] += len(a.data_validations.dataValidation)
        for row in a:
            for cell in row:
                actual = b[cell.coordinate]
                if cell.coordinate not in allowed:
                    assert actual.value == cell.value, (sheet, cell.coordinate, "unrequested content")
                if cell.data_type == "f":
                    assert actual.value == cell.value, (sheet, cell.coordinate, "formula")
                    checks["original_formulas"] += 1
                if cell.fill.fgColor.rgb == "FFDDF0DF":
                    assert actual.value == cell.value and actual.style_id == cell.style_id, (sheet, cell.coordinate, "lecturer")
                    checks["lecturer_cells"] += 1
                date_format = next((u for u in by_sheet.get(sheet, []) if u["cell"] == cell.coordinate and u.get("format_from")), None)
                if date_format:
                    expected_style = deepcopy(cell._style)
                    expected_style.numFmtId = a[date_format["format_from"]]._style.numFmtId
                    assert actual._style == expected_style, (sheet, cell.coordinate, "date format only")
                else:
                    assert actual.style_id == cell.style_id, (sheet, cell.coordinate, "style")
    cached = openpyxl.load_workbook(target, data_only=True)
    checks["client_count"] = cached["Client Q&A"]["C16"].value
    checks["client_limit_check"] = cached["Client Q&A"]["C18"].value
    assert checks["client_limit_check"] == "OK"
    checks["requirements_FR"] = cached["Requirements List"]["C63"].value
    checks["requirements_NFR"] = cached["Requirements List"]["D63"].value
    assert (checks["requirements_FR"], checks["requirements_NFR"]) == (40, 12)
    checks["formula_text_preserved"] = True
    checks["lecturer_values_and_colors_preserved"] = True
    checks["new_assumption_date_formats_only"] = 8
    target.replace(BASE / "workbook" / NAME)
    (BASE / "notes/P1_workbook_preservation.json").write_text(json.dumps(checks, indent=2), encoding="utf-8")
    print(json.dumps(checks, indent=2))


if __name__ == "__main__":
    main()
